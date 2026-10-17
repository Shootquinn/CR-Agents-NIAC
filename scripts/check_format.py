"""check_format.py - NASA NIAC proposal format checker.

Usage: python check_format.py FILE.docx [--chart CHART.docx] [--pdf FILE.pdf]
           [--ai-log AI_LOG.md]
       python check_format.py --chart-only CHART.docx   (flag may come before or after the file)

Prints one line per check:  PASS|FAIL|WARN name value limit
Exit code 1 if any FAIL, else 0.

Definitions (solicitation: at most 15 characters per horizontal inch and
5.5 lines per vertical inch, at least 1-inch margins, US Letter).

cpi = characters including spaces per body line / text width in inches,
      averaged over all body paragraphs. Body paragraphs are the non-empty
      paragraphs whose style is not a Heading or Title style. Lines per
      paragraph are ESTIMATED, not rendered: each word is measured with a
      Times-like per-character width table (units of em, below), scaled by
      the run font size, and words are wrapped greedily into the text width
      (page width minus left and right margins). The table is normalised at
      load time so that ordinary English prose averages 0.40 em per
      character including spaces, which is exactly 15.0 cpi at 12 pt Times
      New Roman (72 pt / (0.40 * 12 pt) = 15.0). Other sizes scale
      linearly (11 pt = 16.4 cpi, 10 pt = 18.0 cpi). Partial last lines
      count as lines, so a paragraph set at 12 pt reads slightly under 15.
      Result = total characters / total estimated lines / text width.
      Non-Times fonts are not re-measured; the same table is used.
lpi = lines per page / text height in inches. Each estimated line is
      font size * 1.15 pt for single spacing (Times line height), the
      exact value for "exactly" spacing, font size * 1.15 * multiple for
      multiple spacing; paragraph space before and after is added per
      paragraph. lpi = total lines / total vertical points * 72.
Both values are approximations for a Word render; the PDF is the final
authority.
"""
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # Windows cp1252 consoles
import re
import sys

DASHES = "\u2014\u2013"
LINE_FACTOR = 1.15

# Times-like advance widths in em (approximate Times New Roman).
_W = {" ": .25, "a": .444, "b": .5, "c": .444, "d": .5, "e": .444, "f": .333,
      "g": .5, "h": .5, "i": .278, "j": .278, "k": .5, "l": .278, "m": .778,
      "n": .5, "o": .5, "p": .5, "q": .5, "r": .333, "s": .389, "t": .278,
      "u": .5, "v": .5, "w": .722, "x": .5, "y": .5, "z": .444,
      ",": .25, ".": .25, ";": .278, ":": .278, "-": .333, "(": .333,
      ")": .333, "'": .18, '"': .408, "/": .278, "%": .833, "&": .778}
for _c in "0123456789":
    _W[_c] = .5
_UP = {"A": .722, "B": .667, "C": .667, "D": .722, "E": .611, "F": .556,
       "G": .722, "H": .722, "I": .333, "J": .389, "K": .722, "L": .611,
       "M": .889, "N": .722, "O": .722, "P": .556, "Q": .722, "R": .667,
       "S": .556, "T": .611, "U": .722, "V": .722, "W": .944, "X": .722,
       "Y": .722, "Z": .611}
_W.update(_UP)
_REF = ("The quick brown fox jumps over the lazy dog while the system "
        "measures density and the reviewers read each page with care. "
        "A balloon floats in the upper atmosphere of the planet and "
        "samples the clouds for signs of life, and the mission returns "
        "data to Earth over a long and steady cruise.")
_RAW = sum(_W.get(c, .5) for c in _REF) / len(_REF)
_K = 0.40 / _RAW


def cw(c):
    return _W.get(c, _W.get(c.lower(), .5)) * _K


def word_w(word, size):
    return sum(cw(c) for c in word) * size / 72.0  # inches


def est_lines(text, size, width_in):
    space = cw(" ") * size / 72.0
    lines, cur = 1, 0.0
    for w in text.split(" "):
        ww = word_w(w, size)
        if cur == 0.0:
            cur = ww
        elif cur + space + ww <= width_in:
            cur += space + ww
        else:
            lines += 1
            cur = ww
    return lines


results = []


def emit(status, name, value, limit):
    results.append(status)
    print("%s %s %s %s" % (status, name, value, limit))


def emu_in(v):
    return None if v is None else v / 914400.0


def style_attr(style, getter):
    while style is not None:
        v = getter(style)
        if v is not None:
            return v
        style = style.base_style
    return None


def body_font(doc):
    normal = doc.styles["Normal"]
    name = style_attr(normal, lambda s: s.font.name)
    size = style_attr(normal, lambda s: s.font.size)
    if name is None or size is None:
        rpr = doc.styles.element.find(
            ".//{*}docDefaults/{*}rPrDefault/{*}rPr")
        if rpr is not None:
            if name is None:
                rf = rpr.find("{*}rFonts")
                if rf is not None:
                    name = next((v for k, v in rf.attrib.items()
                                 if k.endswith("ascii")), None)
            if size is None:
                sz = rpr.find("{*}sz")
                if sz is not None:
                    size = int(next(iter(sz.attrib.values()))) * 6350
    if name is None:
        from collections import Counter
        c = Counter(r.font.name for pp in doc.paragraphs for r in pp.runs
                    if r.font.name)
        if c:
            name = c.most_common(1)[0][0] + " (runs)"
        else:
            for rel in doc.part.rels.values():
                if rel.reltype.endswith("/theme"):
                    m = re.search(rb"<a:minorFont>\s*<a:latin typeface=\"([^\"]+)\"",
                                  rel.target_part.blob)
                    if m:
                        name = m.group(1).decode() + " (theme)"
    return name, (size / 12700.0 if size else None)


def para_metrics(p, dsize):
    pf = p.paragraph_format
    st = p.style
    size = None
    for r in p.runs:
        if r.font.size:
            size = r.font.size.pt
            break
    if size is None:
        s = style_attr(st, lambda s: s.font.size)
        size = s.pt if s else dsize
    ls = pf.line_spacing
    if ls is None:
        ls = style_attr(st, lambda s: s.paragraph_format.line_spacing)
    if ls is None:
        line_h = size * LINE_FACTOR
    elif hasattr(ls, "pt"):
        line_h = ls.pt
    else:
        line_h = size * LINE_FACTOR * float(ls)
    sb = pf.space_before
    if sb is None:
        sb = style_attr(st, lambda s: s.paragraph_format.space_before)
    sa = pf.space_after
    if sa is None:
        sa = style_attr(st, lambda s: s.paragraph_format.space_after)
    return size, line_h, (sb.pt if sb else 0.0), (sa.pt if sa else 0.0)


def is_body(p):
    n = ((p.style.name if p.style is not None else "") or "").lower()
    return p.text.strip() and not (n.startswith("heading") or n == "title")


def all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.extend(p.text for p in c.paragraphs)
    return "\n".join(parts)


def page_dims(doc):
    sec = doc.sections[0]
    return sec, emu_in(sec.page_width), emu_in(sec.page_height)


def check_body(path):
    import docx
    doc = docx.Document(path)
    sec, w, h = page_dims(doc)
    ok = abs(w - 8.5) < 0.01 and abs(h - 11.0) < 0.01
    emit("PASS" if ok else "FAIL", "page_size", "%.2fx%.2f" % (w, h),
         "8.5x11")
    for nm in ("left", "right", "top", "bottom"):
        v = emu_in(getattr(sec, nm + "_margin"))
        emit("PASS" if v >= 0.995 else "FAIL", "margin_" + nm,
             "%.2f" % v, ">=1.0")
    fname, fsize = body_font(doc)
    text_w = w - emu_in(sec.left_margin) - emu_in(sec.right_margin)
    chars = lines = all_lines = 0
    height = 0.0
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        size, line_h, sb, sa = para_metrics(p, fsize or 12.0)
        n = est_lines(p.text, size, text_w)
        if is_body(p):
            chars += len(p.text)
            lines += n
        all_lines += n
        height += n * line_h + sb + sa
    cpi = chars / float(lines) / text_w if lines else 0.0
    lpi = all_lines / (height / 72.0) if height else 0.0
    emit("PASS" if cpi <= 15.0 else "FAIL", "cpi", "%.2f" % cpi, "<=15")
    emit("PASS" if lpi <= 5.5 else "FAIL", "lpi", "%.2f" % lpi, "<=5.5")
    nd = sum(all_text(doc).count(c) for c in DASHES)
    emit("PASS" if nd == 0 else "FAIL", "dash_count", nd, "0")
    emit("PASS" if fname and fsize else "WARN", "body_font",
         "%s %s pt" % (fname, fsize), "read from styles")
    low = imgs = 0
    for shp in doc.inline_shapes:
        try:
            rid = shp._inline.graphic.graphicData.pic.blipFill.blip.embed
            img = doc.part.related_parts[rid].image
            imgs += 1
            dpi = img.px_width / emu_in(shp.width)
            if dpi < 200:
                low += 1
                emit("WARN", "image_dpi", "%.0f" % dpi, ">=200")
        except Exception:
            continue
    if low == 0:
        emit("PASS", "image_dpi", "%d images" % imgs, ">=200")
    check_styles(doc)
    check_figures(doc, path)


STYLE_WANT = (("Normal", 12.0), ("Title", 12.0), ("Heading 1", 12.0),
              ("Heading 2", 12.0), ("Caption", 10.0))
TNR = "times new roman"


def default_para_spacing(doc):
    """(before, after, line, rule) from docDefaults pPrDefault, in points
    and multiples; None for a missing value."""
    sp = doc.styles.element.find(".//{*}docDefaults/{*}pPrDefault/{*}pPr/{*}spacing")
    out = {"before": None, "after": None, "line": None, "rule": "auto"}
    if sp is not None:
        for k in ("before", "after", "line"):
            v = sp.get(W_NS + k)
            if v is not None:
                out[k] = int(v)
        out["rule"] = sp.get(W_NS + "lineRule") or "auto"
    return out


def style_readback(doc, name):
    try:
        st = doc.styles[name]
    except KeyError:
        return None
    fname = style_attr(st, lambda s: s.font.name)
    if fname is None:
        fname = body_font(doc)[0]
    size = style_attr(st, lambda s: s.font.size)
    dd = default_para_spacing(doc)
    sb = style_attr(st, lambda s: s.paragraph_format.space_before)
    sa = style_attr(st, lambda s: s.paragraph_format.space_after)
    ls = style_attr(st, lambda s: s.paragraph_format.line_spacing)
    before = sb.pt if sb is not None else (
        dd["before"] / 20.0 if dd["before"] is not None else 0.0)
    after = sa.pt if sa is not None else (
        dd["after"] / 20.0 if dd["after"] is not None else 0.0)
    if ls is None:
        ls = (dd["line"] / 240.0 if dd["line"] is not None
              and dd["rule"] == "auto" else 1.0)
    single = not hasattr(ls, "pt") and abs(float(ls) - 1.0) < 0.001
    lval = "exactly %.1fpt" % ls.pt if hasattr(ls, "pt") else "%.2f" % float(ls)
    return fname, (size.pt if size else None), before, after, single, lval


def check_styles(doc):
    for name, want in STYLE_WANT:
        key = "style_" + name.lower().replace(" ", "_")
        r = style_readback(doc, name)
        if r is None:
            emit("FAIL", key, "missing", "TNR %g pt 0/4 single" % want)
            continue
        fname, size, before, after, single, lval = r
        ok = ((fname or "").lower() == TNR and size == want
              and abs(before) < 0.01 and abs(after - 4.0) < 0.01 and single)
        emit("PASS" if ok else "FAIL", key,
             "%s %s pt before %.1f after %.1f line %s" % (
                 fname, size, before, after, lval),
             "TNR %g pt 0/4 single" % want)


def check_figures(doc, path):
    import zipfile
    from lxml import etree
    root = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    anchors = len(list(root.iter("{*}anchor")))
    inline = len(list(root.iter("{*}inline")))
    emit("PASS" if anchors == 0 else "FAIL", "figure_inline",
         "%d inline %d floating" % (inline, anchors), "0 floating")
    txt = all_text(doc)
    miss = txt.count("[FIGURE")
    emit("PASS" if miss == 0 else "FAIL", "image_missing", miss, "0")
    ip = txt.count("Image prompt:")
    emit("PASS" if ip == 0 else "FAIL", "image_prompt_text", ip, "0")


def pdf_info(path):
    """List of (text, width_pt, height_pt) per page, or None when no PDF
    library imports."""
    try:
        import pypdf
        out = []
        for pg in pypdf.PdfReader(path).pages:
            w, h = float(pg.mediabox.width), float(pg.mediabox.height)
            if (pg.get("/Rotate") or 0) % 180 == 90:
                w, h = h, w
            out.append((pg.extract_text() or "", w, h))
        return out
    except ImportError:
        pass
    try:
        import fitz
        return [(pg.get_text(), pg.rect.width, pg.rect.height)
                for pg in fitz.open(path)]
    except ImportError:
        return None


def pdf_pages(path):
    info = pdf_info(path)
    return None if info is None else [t for t, _, _ in info]


AI_MODEL = re.compile(
    r"(?<![A-Za-z])(Claude|Haiku|Opus|Sonnet|ChatGPT|GPT-?[0-9]*|Gemini|"
    r"Copilot|Llama|Mistral|Grok|DeepSeek|Qwen|Imagen|Midjourney|DALL-?E|"
    r"Firefly|Stable Diffusion|Nano Banana)(?![A-Za-z])")
MONTHS = ("january february march april may june july august september "
          "october november december").split()
_MN = "|".join(m.capitalize() for m in MONTHS)
DATE_RE = re.compile(
    r"(?<![0-9])([0-9]{4})-([0-9]{2})-([0-9]{2})(?![0-9])"
    r"|(" + _MN + r")\s+([0-9]{1,2}),?\s+([0-9]{4})(?![0-9])"
    r"|(?<![0-9])([0-9]{1,2})\s+(" + _MN + r")\s+([0-9]{4})(?![0-9])")



def iso_dates(text):
    """(iso date, start, end) for every date in text."""
    out = []
    for m in DATE_RE.finditer(text):
        g = m.groups()
        if g[0]:
            iso = "%s-%s-%s" % (g[0], g[1], g[2])
        elif g[3]:
            iso = "%s-%02d-%02d" % (g[5], MONTHS.index(g[3].lower()) + 1,
                                    int(g[4]))
        else:
            iso = "%s-%02d-%02d" % (g[8], MONTHS.index(g[7].lower()) + 1,
                                    int(g[6]))
        out.append((iso, m.start(), m.end()))
    return out


def ai_citation_dates(text):
    """ISO dates that sit within 200 characters of an AI model name."""
    flat = re.sub(r"\s+", " ", text)
    flat = re.sub("([0-9]-) ([0-9])", lambda m: m.group(1) + m.group(2), flat)
    found = []
    for iso, a, b in iso_dates(flat):
        if AI_MODEL.search(flat[max(0, a - 200):b + 200]):
            found.append(iso)
    return found


def log_rows_with(path, iso):
    try:
        rows = open(path, encoding="utf-8").read().splitlines()
    except OSError:
        return 0
    return sum(1 for r in rows if r.lstrip().startswith("|") and iso in r)


def check_pdf(path, ai_log=None):
    import os
    info = pdf_info(path)
    if info is None:
        print("RENDER TOOL MISSING")
        emit("FAIL", "pdf_chart_page1_landscape", "n/a", "landscape")
        emit("FAIL", "pdf_body_pages", "n/a", "<=3")
        emit("FAIL", "pdf_body_ends_by_page4", "n/a", "<=4")
        emit("FAIL", "pdf_references_page", "n/a", ">=5")
        return
    pages = [t for t, _, _ in info]
    n = len(pages)
    w, h = (info[0][1], info[0][2]) if info else (0, 0)
    chart = 1 if w > h else 0
    emit("PASS" if chart else "FAIL", "pdf_chart_page1_landscape",
         "%.0fx%.0f pt" % (w, h), "width > height")
    ref = None
    for i in range(chart, n):
        if re.search(r"^\s*references\s*$", pages[i], re.I | re.M):
            ref = i + 1
            break
    last = (ref - 1) if ref else n
    body = last - chart
    emit("PASS" if body <= 3 else "FAIL", "pdf_body_pages", body, "<=3")
    emit("PASS" if last <= chart + 3 else "FAIL", "pdf_body_ends_by_page4",
         last, "<=%d" % (chart + 3))
    if ref is None:
        emit("PASS", "pdf_references_page", "none", ">=5 if present")
    else:
        emit("PASS" if ref >= 5 else "FAIL", "pdf_references_page", ref,
             ">=5")
    mb = os.path.getsize(path) / 1048576.0
    emit("PASS" if mb < 20 else "FAIL", "pdf_size_mb", "%.2f" % mb, "<20")
    text = chr(10).join(pages)
    ip = text.count("Image prompt:")
    emit("PASS" if ip == 0 else "FAIL", "pdf_image_prompt", ip, "0")
    dates = ai_citation_dates(text)
    if not dates:
        emit("FAIL", "pdf_ai_citation", "none", "AI model line with date")
    elif ai_log:
        hit = [d for d in dates if log_rows_with(ai_log, d)]
        emit("PASS" if hit else "FAIL", "pdf_ai_citation",
             "%s in log" % hit[0] if hit else "%s not in log" % dates[0],
             "date in AI log row")
    else:
        emit("PASS", "pdf_ai_citation", dates[0], "AI model line with date")


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PIC_NS = "{http://schemas.openxmlformats.org/drawingml/2006/picture}"


def in_fallback(el):
    """True when el sits inside mc:Fallback (the legacy copy of a text box)."""
    for a in el.iterancestors():
        if a.tag.endswith("}Fallback"):
            return True
    return False


def theme_fonts(zf):
    """Return (major, minor) latin typefaces from word/theme/theme1.xml."""
    from lxml import etree
    major = minor = None
    for n in zf.namelist():
        if n.startswith("word/theme/") and n.endswith(".xml"):
            root = etree.fromstring(zf.read(n))
            mj = root.find(".//" + A_NS + "majorFont/" + A_NS + "latin")
            mn = root.find(".//" + A_NS + "minorFont/" + A_NS + "latin")
            if mj is not None:
                major = mj.get("typeface")
            if mn is not None:
                minor = mn.get("typeface")
            break
    return major, minor


class StyleMap:
    """Style chain resolver reading styles.xml directly, so theme fonts
    (asciiTheme) and docDefaults are honoured."""

    def __init__(self, zf):
        from lxml import etree
        self.major, self.minor = theme_fonts(zf)
        root = etree.fromstring(zf.read("word/styles.xml"))
        self.styles = {}
        self.default_par = None
        for st in root.findall(W_NS + "style"):
            sid = st.get(W_NS + "styleId")
            self.styles[sid] = st
            if (st.get(W_NS + "type") == "paragraph"
                    and st.get(W_NS + "default") in ("1", "true")):
                self.default_par = sid
        dd = root.find(W_NS + "docDefaults/" + W_NS + "rPrDefault/"
                       + W_NS + "rPr")
        self.dname, self.dsize = self.rpr_font(dd)

    def theme_name(self, ref):
        if ref.startswith("major"):
            return self.major
        return self.minor

    def rpr_font(self, rpr):
        name = size = None
        if rpr is None:
            return name, size
        rf = rpr.find(W_NS + "rFonts")
        if rf is not None:
            name = rf.get(W_NS + "ascii")
            if name is None and rf.get(W_NS + "asciiTheme"):
                name = self.theme_name(rf.get(W_NS + "asciiTheme"))
            if name is None:
                name = rf.get(W_NS + "hAnsi")
            if name is None and rf.get(W_NS + "hAnsiTheme"):
                name = self.theme_name(rf.get(W_NS + "hAnsiTheme"))
        sz = rpr.find(W_NS + "sz")
        if sz is not None:
            size = int(sz.get(W_NS + "val")) / 2.0
        return name, size

    def chain(self, sid):
        seen = set()
        while sid and sid in self.styles and sid not in seen:
            seen.add(sid)
            st = self.styles[sid]
            yield st
            b = st.find(W_NS + "basedOn")
            sid = b.get(W_NS + "val") if b is not None else None

    def resolve(self, name, size, sid):
        for st in self.chain(sid):
            if name is not None and size is not None:
                break
            n, s = self.rpr_font(st.find(W_NS + "rPr"))
            name = name or n
            size = size or s
        return name, size

    def run_font(self, r_el):
        """Resolve (font, size pt): run rPr, character style, paragraph
        style, default paragraph style, docDefaults."""
        rpr = r_el.find(W_NS + "rPr")
        name, size = self.rpr_font(rpr)
        if rpr is not None and rpr.find(W_NS + "rStyle") is not None:
            name, size = self.resolve(
                name, size, rpr.find(W_NS + "rStyle").get(W_NS + "val"))
        p = r_el.getparent()
        while p is not None and p.tag != W_NS + "p":
            p = p.getparent()
        sid = None
        if p is not None:
            ppr = p.find(W_NS + "pPr")
            if ppr is not None and ppr.find(W_NS + "pStyle") is not None:
                sid = ppr.find(W_NS + "pStyle").get(W_NS + "val")
        name, size = self.resolve(name, size, sid or self.default_par)
        return name or self.dname, size or self.dsize


def run_text(r_el):
    out = []
    for c in r_el:
        if c.tag == W_NS + "t":
            out.append(c.text or "")
        elif c.tag in (W_NS + "tab", W_NS + "br", W_NS + "cr"):
            out.append(" ")
    return "".join(out)


def chart_runs(root):
    """Yield every w:r outside mc:Fallback that has its own text (blank
    runs included, since they carry the spaces between words). Each text
    box is counted once because only the mc:Choice copy is read, and run
    text comes from the run's own children (not its descendants, which
    include the text of any text box nested in the run)."""
    for r in root.iter(W_NS + "r"):
        if in_fallback(r):
            continue
        if run_text(r):
            yield r


def chart_pictures(root, zf, rels):
    """Return a list of (width_px, width_in), one per picture outside
    mc:Fallback."""
    from docx.image.image import Image
    out = []
    for pic in root.iter(PIC_NS + "pic"):
        if in_fallback(pic):
            continue
        blip = pic.find(".//" + A_NS + "blip")
        ext = pic.find(PIC_NS + "spPr/" + A_NS + "xfrm/" + A_NS + "ext")
        if blip is None:
            continue
        target = rels.get(blip.get(R_NS + "embed"))
        px = wid = None
        try:
            px = Image.from_blob(zf.read("word/" + target)).px_width
        except Exception:
            pass
        if ext is not None and ext.get("cx"):
            wid = int(ext.get("cx")) / 914400.0
        out.append((px, wid))
    return out


def doc_rels(zf):
    from lxml import etree
    root = etree.fromstring(zf.read("word/_rels/document.xml.rels"))
    return {e.get("Id"): e.get("Target") for e in root}


def box_geometry(box):
    """Return (name, width_emu, height_emu, lIns, tIns, rIns, bIns) for the
    shape that holds txbxContent `box`, or None."""
    WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    wsp = next((a for a in box.iterancestors() if a.tag.endswith("}wsp")),
               None)
    if wsp is None:
        return None
    ins = [91440, 45720, 91440, 45720]
    bp = wsp.find("{*}bodyPr")
    if bp is not None:
        for i, k in enumerate(("lIns", "tIns", "rIns", "bIns")):
            if bp.get(k):
                ins[i] = int(bp.get(k))
    grp = next((a for a in box.iterancestors() if a.tag.endswith("}wgp")),
               None)
    if grp is None:
        anchor = next(a for a in box.iterancestors()
                      if a.tag.endswith("}anchor"))
        e = anchor.find(WP + "extent")
        return (anchor.find(WP + "docPr").get("name"), int(e.get("cx")),
                int(e.get("cy"))) + tuple(ins)
    xf = grp.find("{*}grpSpPr/" + A_NS + "xfrm")
    ext, chext = xf.find(A_NS + "ext"), xf.find(A_NS + "chExt")
    sx = int(ext.get("cx")) / float(chext.get("cx"))
    sy = int(ext.get("cy")) / float(chext.get("cy"))
    e = wsp.find("{*}spPr/" + A_NS + "xfrm/" + A_NS + "ext")
    nm = wsp.find("{*}cNvPr")
    return (nm.get("name") if nm is not None else "box",
            int(e.get("cx")) * sx, int(e.get("cy")) * sy) + tuple(ins)


def box_text_height(box, sm, width_pt):
    """Estimated text height in points: Arial average char width 0.5 em,
    line height 1.15 x size, plus paragraph spacing."""
    total = 0.0
    for p in box.findall(W_NS + "p"):
        runs = [r for r in p.iter(W_NS + "r")]
        text = "".join(run_text(r) for r in runs)
        sizes = [sm.run_font(r)[1] or 11.0 for r in runs] or [11.0]
        size = max(sizes)
        ppr = p.find(W_NS + "pPr")
        indent = 0.0
        before = after = 0.0
        if ppr is not None:
            ind = ppr.find(W_NS + "ind")
            if ind is not None and ind.get(W_NS + "left"):
                indent = int(ind.get(W_NS + "left")) / 20.0
            elif ppr.find(W_NS + "numPr") is not None:
                indent = 36.0
            sp = ppr.find(W_NS + "spacing")
            if sp is not None:
                before = int(sp.get(W_NS + "before") or 0) / 20.0
                after = int(sp.get(W_NS + "after") or 0) / 20.0
        avail = max(width_pt - indent, size)
        cw = 0.5 * size
        lines, cur = 1, 0.0
        for wd in text.split():
            ww = len(wd) * cw
            if cur == 0.0:
                cur = ww
            elif cur + cw + ww <= avail:
                cur += cw + ww
            else:
                lines += 1
                cur = ww
        total += lines * 1.15 * size + before + after
    return total


def check_box_overflow(root, sm):
    over, n = [], 0
    for box in root.iter(W_NS + "txbxContent"):
        if in_fallback(box):
            continue
        g = box_geometry(box)
        if g is None:
            continue
        name, wemu, hemu, li, ti, ri, bi = g
        n += 1
        have = (hemu - ti - bi) / 12700.0
        need = box_text_height(box, sm, (wemu - li - ri) / 12700.0)
        if need > have:
            over.append("%s need %.2fin have %.2fin" % (
                name, need / 72.0, have / 72.0))
    emit("FAIL" if over else "PASS", "chart_box_overflow",
         "; ".join(over) if over else "0 of %d boxes over" % n,
         "text height <= box height")


def check_chart(path):
    import zipfile
    import docx
    from lxml import etree
    doc = docx.Document(path)
    sec, w, h = page_dims(doc)
    ok = abs(w - 11.0) < 0.01 and abs(h - 8.5) < 0.01
    emit("PASS" if ok else "FAIL", "chart_landscape", "%.2fx%.2f" % (w, h),
         "11x8.5")
    for nm in ("left", "right", "top", "bottom"):
        v = emu_in(getattr(sec, nm + "_margin"))
        emit("PASS" if v >= 0.995 else "FAIL", "chart_margin_" + nm,
             "%.2f" % v, ">=1.0")
    zf = zipfile.ZipFile(path)
    sm = StyleMap(zf)
    root = etree.fromstring(zf.read("word/document.xml"))
    small = []
    paras = {}
    for r in chart_runs(root):
        p = r.getparent()
        while p is not None and p.tag != W_NS + "p":
            p = p.getparent()
        paras.setdefault(p, []).append(run_text(r))
        if not run_text(r).strip():
            continue
        n, s = sm.run_font(r)
        if s is None or s < 12 or (n or "").lower() != "arial":
            small.append("%s %s" % (n, s))
    emit("PASS" if not small else "FAIL", "chart_font_arial12",
         "%d runs under Arial 12 (%s)" % (len(small),
                                         small[0] if small else "-"), "0")
    txt = " ".join("".join(v) for v in paras.values())
    for wd in ("schedule", "team", "cost"):
        n = len(re.findall(r"\b%s\b" % wd, txt, re.I))
        emit("PASS" if n == 0 else "FAIL", "chart_word_" + wd, n, "0")
    wc = len(txt.split())
    emit("PASS" if 200 <= wc <= 360 else "FAIL", "chart_word_count", wc,
         "200-360")
    nref = len(re.findall(r"\bsee (section|figure)\b", txt, re.I))
    emit("PASS" if nref == 0 else "FAIL", "chart_see_refs", nref, "0")
    nd = sum(txt.count(c) for c in DASHES)
    emit("PASS" if nd == 0 else "FAIL", "chart_dash_count", nd, "0")
    check_box_overflow(root, sm)
    pics = chart_pictures(root, zf, doc_rels(zf))
    emit("PASS" if 1 <= len(pics) <= 3 else "FAIL", "chart_images",
         len(pics), "1-3")
    low = [px / wid for px, wid in pics if px and wid and px / wid < 200]
    if low:
        emit("WARN", "chart_image_dpi", "%.0f" % min(low), ">=200")
    else:
        dpis = [px / wid for px, wid in pics if px and wid]
        emit("PASS", "chart_image_dpi", "%d images min %s dpi" % (
            len(pics), "%.0f" % min(dpis) if dpis else "n/a"), ">=200")


def main(argv):
    args = argv[1:]
    valued = ("--chart", "--pdf", "--ai-log")
    pos, flags, i = [], {}, 0
    while i < len(args):
        a = args[i]
        if a in valued and i + 1 < len(args):
            flags[a] = args[i + 1]
            i += 2
            continue
        if a.startswith("--"):
            flags[a] = True
        else:
            pos.append(a)
        i += 1
    if "--chart-only" in flags:
        # works in any argument position: the one non-flag argument is the chart
        # (a value given with --chart-only itself is also accepted)
        target = pos[0] if pos else (
            flags["--chart-only"] if isinstance(flags["--chart-only"], str)
            else None)
        if not target:
            print("usage: check_format.py --chart-only CHART.docx")
            return 2
        check_chart(target)
        return 1 if "FAIL" in results else 0
    if not pos:
        print("usage: check_format.py FILE.docx [--chart CHART.docx] "
              "[--pdf FILE.pdf]\n       check_format.py --chart-only "
              "CHART.docx")
        return 2
    check_body(pos[0])
    if flags.get("--pdf"):
        check_pdf(flags["--pdf"], flags.get("--ai-log"))
    if flags.get("--chart"):
        check_chart(flags["--chart"])
    return 1 if "FAIL" in results else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
