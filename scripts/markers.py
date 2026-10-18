"""markers.py - the one definition of process text that must never reach
the proposal. build_proposal_docx.py, check_format.py and the Manager use it.

Usage: python scripts/markers.py FILE   (.md: the ## Draft or ## Edited draft
       section; .docx: all paragraphs; .pdf: all page text)

Kinds:
  marker  [SOURCE NEEDED], [REFERENCE NEEDED], [FIGURE UNVERIFIED], UNVERIFIED,
          NOT-READ, with or without brackets
  path    cr_scratch, a drive-letter path (C:/ or C:\\), /Users/ or \\Users\\
  file    a file name ending in .md, .txt or .py
  line    "line 37", "lines 12-14", "L34:" (a file line reference)
  step    "step 1", "steps 2 and 3" (only NIAC Step A and Step B belong in
          the body; Phase I activities are named, not numbered)
Prints one line per hit, `HIT kind: text`, then `MARKERS n`. Exit 0.
"""
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # Windows cp1252 consoles

BRACKET_RE = re.compile(r"\s*\[(?:SOURCE NEEDED|REFERENCE NEEDED|FIGURE UNVERIFIED|UNVERIFIED|NOT-READ)[^\]]*\]")
KINDS = (
    ("marker", re.compile(r"SOURCE NEEDED|REFERENCE NEEDED|FIGURE UNVERIFIED|\bUNVERIFIED\b|NOT-READ")),
    ("path", re.compile(r"cr_scratch|\b[A-Za-z]:[\\/]\S*|[\\/]Users[\\/]\S*")),
    ("file", re.compile(r"\b[\w-]+\.(?:md|txt|py)\b")),
    ("line", re.compile(r"\blines? \d+(?:\s*(?:-|to|and)\s*\d+)?\b|\bL\d+:", re.I)),
    ("step", re.compile(r"\bsteps? \d+(?:\s*(?:,|and|to|-)\s*\d+)*\b", re.I)),
)
PATH_LINE_RE = re.compile(r"^\s*`?(?:[A-Za-z]:[\\/]|[\\/]|\.{0,2}[\\/]?cr_scratch[\\/]).*\.(?:md|txt|docx|pdf|py)`?\s*$")


def find(text):
    """List of (kind, matched text) for every hit in text."""
    return [(k, m.group(0)) for k, rx in KINDS for m in rx.finditer(text)]


def strip_brackets(line):
    """(line without bracketed markers, number removed)."""
    return BRACKET_RE.sub("", line), len(BRACKET_RE.findall(line))


def is_path_line(line):
    """True for a line that holds nothing but a file path."""
    return bool(PATH_LINE_RE.match(line))


def draft_text(path):
    out, on = [], False
    for ln in open(path, encoding="utf-8-sig").read().splitlines():
        if ln.startswith("## "):
            if on:
                break
            on = ln.strip().lower() in ("## draft", "## edited draft")
            continue
        if on:
            out.append(ln)
    return "\n".join(out)


def file_text(path):
    low = path.lower()
    if low.endswith(".docx"):
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    parts.extend(p.text for p in c.paragraphs)
        return "\n".join(parts)
    if low.endswith(".pdf"):
        import pypdf
        return "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(path).pages)
    return draft_text(path)


def main():
    if len(sys.argv) != 2 or not os.path.exists(sys.argv[1]):
        print("USAGE: python scripts/markers.py FILE (.md draft, .docx or .pdf)")
        return 2
    hits = find(file_text(sys.argv[1]))
    for k, t in hits:
        print("HIT %s: %s" % (k, t))
    print("MARKERS %d" % len(hits))
    return 0


if __name__ == "__main__":
    sys.exit(main())
