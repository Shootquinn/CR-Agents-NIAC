"""build_proposal_docx.py - build the proposal body docx from a Writer draft.

Usage: build_proposal_docx.py --draft cr_scratch/stepN_writer_draft.md
         --template templates/proposal_template.docx --figures figures
         --out cr_scratch/proposal_body.docx
         [--sources-to-get cr_scratch/sources_to_get.md]

Draft: the text under the `## Draft` heading, up to the next `## ` heading.
  1. Plain paragraphs separated by blank lines. Run-in labels `**Label.**`
     become bold.
  2. `# text` is the Title; `### text` is Heading 1; `#### text` is Heading 2.
  3. `Figure N: caption.` is a figure. An optional `Image prompt: ...` line
     after it is removed from the document and written to FIGURES/prompts/NN.txt.
  4. Image for figure N is FIGURES/NN.png or NN.jpg. With an image: the picture
     inline (at most text width and 2.5 in tall), then the caption in the
     Caption style. Without an image: the line `[FIGURE NN IMAGE MISSING]`, the
     caption, and the printed line `MISSING IMAGE NN` (exit 0).
  5. A line `References` starts the reference list on a new page (References
     style). Each following non-empty line is one entry.
  6. --sources-to-get FILE (optional): its numbered or bulleted lines are
     appended after the references under the line SOURCES TO GET BEFORE THE
     NEXT REVISION. Prints `SOURCES TO GET n`. A missing or empty file adds
     nothing.
Every TEMPLATE NOTE paragraph is removed. Exit 0 on success, 2 on a bad input.
"""
import argparse
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # Windows cp1252 consoles

import docx
from docx.shared import Inches

MAX_H = 2.5
FIG_RE = re.compile(r"^Figure\s+(\d+)\s*:\s*(.*)$")
PROMPT_RE = re.compile(r"^Image prompt\s*:\s*(.*)$", re.I)
REF_RE = re.compile(r"^(?:#+\s*)?(?:\*\*)?References(?:\*\*)?\s*$", re.I)
BOLD_RE = re.compile(r"(\*\*[^*]+\*\*)")


def draft_lines(path):
    out, on = [], False
    for ln in open(path, encoding="utf-8").read().splitlines():
        if ln.startswith("## "):
            if on:
                break
            on = ln.strip().lower() == "## draft"
            continue
        if on:
            out.append(ln.rstrip())
    return out


def add_runs(p, text):
    for part in BOLD_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        else:
            p.add_run(part)


def find_image(figdir, n):
    for ext in ("png", "jpg", "jpeg", "PNG", "JPG", "JPEG"):
        f = os.path.join(figdir, "%02d.%s" % (n, ext))
        if os.path.isfile(f):
            return f
    return None


def add_picture(doc, path, text_w):
    from docx.image.image import Image
    img = Image.from_file(path)
    w = float(text_w)
    h = w * img.px_height / img.px_width
    if h > MAX_H:
        w, h = MAX_H * img.px_width / img.px_height, MAX_H
    p = doc.add_paragraph(style="Normal")
    p.alignment = 1
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(w), height=Inches(h))


def to_get_lines(path):
    if not path or not os.path.exists(path):
        return []
    out = []
    for ln in open(path, encoding="utf-8").read().splitlines():
        m = re.match(r"^\s*(?:[-*]|\d+\.)\s+(.+)$", ln)
        if m:
            out.append(m.group(1).strip())
    return out


def build(draft, template, figdir, out, to_get=None):
    lines = draft_lines(draft)
    if not any(x.strip() for x in lines):
        print("NO DRAFT TEXT under ## Draft in %s" % draft)
        return 2
    doc = docx.Document(template)
    for p in list(doc.paragraphs):
        if p.text.strip().startswith("TEMPLATE NOTE"):
            p._element.getparent().remove(p._element)
    sec = doc.sections[0]
    text_w = (sec.page_width - sec.left_margin - sec.right_margin) / 914400.0
    prompts = {}
    figs = []
    para, in_refs, last_fig = [], False, None

    def flush():
        if para:
            add_runs(doc.add_paragraph(style="Normal"), " ".join(para))
            para.clear()

    for ln in lines:
        s = ln.strip()
        if not s:
            flush()
            continue
        m = FIG_RE.match(s)
        if m and not in_refs:
            flush()
            n, cap = int(m.group(1)), m.group(2).strip()
            last_fig = n
            img = find_image(figdir, n)
            figs.append((n, img))
            if img:
                add_picture(doc, img, text_w)
            else:
                doc.add_paragraph("[FIGURE %02d IMAGE MISSING]" % n,
                                  style="Normal")
                print("MISSING IMAGE %02d" % n)
            doc.add_paragraph("Figure %d: %s" % (n, cap), style="Caption")
            continue
        m = PROMPT_RE.match(s)
        if m and last_fig is not None:
            flush()
            prompts[last_fig] = m.group(1).strip()
            continue
        if REF_RE.match(s):
            flush()
            h = doc.add_paragraph("References", style="References")
            h.paragraph_format.page_break_before = True
            in_refs = True
            continue
        if in_refs:
            add_runs(doc.add_paragraph(style="Normal"), re.sub(
                r"^(?:[-*]\s+|\d+\.\s+)", "", s))
            continue
        last_fig = None
        if s.startswith("#### "):
            flush()
            doc.add_paragraph(s[5:], style="Heading 2")
        elif s.startswith("### "):
            flush()
            doc.add_paragraph(s[4:], style="Heading 1")
        elif s.startswith("# "):
            flush()
            doc.add_paragraph(s[2:], style="Title")
        else:
            para.append(s)
    flush()
    get = to_get_lines(to_get)
    if get:
        if not in_refs:
            h = doc.add_paragraph("References", style="References")
            h.paragraph_format.page_break_before = True
        doc.add_paragraph("SOURCES TO GET BEFORE THE NEXT REVISION",
                          style="Normal").runs[0].bold = True
        for g in get:
            add_runs(doc.add_paragraph(style="Normal"), g)
    print("SOURCES TO GET %d" % len(get))
    if prompts:
        pdir = os.path.join(figdir, "prompts")
        os.makedirs(pdir, exist_ok=True)
        for n, t in prompts.items():
            with open(os.path.join(pdir, "%02d.txt" % n), "w",
                      encoding="utf-8") as f:
                f.write(t + "\n")
    for n, img in figs:
        print("FIGURE %02d image=%s prompt=%s" % (
            n, os.path.basename(img) if img else "MISSING",
            "figures/prompts/%02d.txt" % n if n in prompts else "none"))
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    doc.save(out)
    print("WROTE %s" % out)
    return 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser(
        description="Build the proposal body docx from a Writer draft. "
                    "Missing figure images print MISSING IMAGE NN (exit 0).")
    ap.add_argument("--draft", required=True)
    ap.add_argument("--template", required=True)
    ap.add_argument("--figures", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--sources-to-get")
    a = ap.parse_args()
    sys.exit(build(a.draft, a.template, a.figures, a.out, a.sources_to_get))
